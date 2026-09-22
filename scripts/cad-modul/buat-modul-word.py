#!/usr/bin/env python3
"""
Membangun Modul-Word (.docx) Pemodelan CAD dari halaman modul HTML-nya.

Skrip ini adalah padanan `scripts/ttl-modul/buat-modul-word.py` untuk mata
kuliah Pemodelan Computer Aided Design; alur, kerangka dokumen, dan seluruh
mesin pengubah HTML/LaTeX → Word-nya sama, hanya sumber isinya yang berbeda.

Sumber isi adalah `Pemodelan-Computer-Aided-Design/Modul/Modul-N.html` (yang
dibangun `bangun-modul-1.py`/`bangun.py` dari `modul_N.py`), sehingga Word dan
HTML tidak pernah menyimpang: pendahuluan (hero), sembilan bagian materi
beserta gambar (SVG dirender ke PNG), persamaan bernomor, kartu, tabel, kotak
catatan/peringatan, ringkasan animasi, cell Python FreeCAD, tugas (10 PG + 5
tugas pemodelan FreeCAD beserta gambar acuannya), forum, dan daftar pustaka.
Hanya `#page-modul`, `#page-tugas`, dan `#page-forum` yang dibaca; tab Setup
FreeCAD dan Setup Python yang hanya ada di Modul 1 tidak ikut ke dokumen.

Kerangka dokumen mengikuti format BOP: sampul, header, footer, dan gaya
diambil dari modul Word yang sudah ada (Modul 1 Sistem Kendali Cerdas, yang
dibuat dari template resmi `Template-Modul-Word-dan-PPT/`), lalu sampulnya
diganti identitas Pemodelan CAD dan seluruh badannya dibangun ulang.
Panel "MODUL INTERAKTIF" dibuat dengan `panel-modul-interaktif-docx.py`
sehingga bentuknya sama dengan modul Word course lain.

"Bahan kajian" dan "Indikator" tanpa berkas RPS
-----------------------------------------------
Teknik Tenaga Listrik mengambil dua butir itu dari `buat-rps.js`. Pemodelan CAD
tidak punya berkas RPS JavaScript, jadi keduanya diturunkan dari sumber resmi
yang memang ada di repositori ini — tidak ada kalimat yang dikarang:

* Bahan kajian ← `Pemodelan-Computer-Aided-Design/Banner/Banner-Pertemuan-P.html`
  (P = pertemuan modul; Pertemuan 8 adalah UTS sehingga Modul 8-14 memakai
  banner Pertemuan 9-15). Banner itu dipakai dosen di LMS dan memuat judul
  topik, satu paragraf cakupan pertemuan, dan deretan chip kata kunci.
  Paragraf cakupannya dipakai apa adanya sebagai bahan kajian, chip-nya sebagai
  kata kunci.
* Indikator ← butir penilaian yang benar-benar dinilai pada modul itu, dibaca
  dari halaman modulnya sendiri: jumlah soal pilihan ganda beserta bobotnya,
  label kelima tugas pemodelan (`compEzDefs`/`compHardDefs`, yaitu deliverable
  .FCStd + angka bacaan yang diperiksa server) beserta poinnya, dan jumlah
  pertanyaan Forum Diskusi. Semuanya terukur dan tertulis di repositori.
* Sub-CPMK dan bobot Tugas/UTS/UAS ← `Attributes/Asesmen-Pemodelan-Computer-Aided-Design.json`.

Pakai:
    python scripts/cad-modul/buat-modul-word.py 5          # satu modul
    python scripts/cad-modul/buat-modul-word.py 1 2 3      # beberapa
    python scripts/cad-modul/buat-modul-word.py --semua    # Modul 1-14
Lalu:
    python scripts/docx-ke-pdf.py Pemodelan-Computer-Aided-Design/Modul-Word/*.docx
    python scripts/cad-modul/pasang-tautan-pdf.py
    python scripts/gabung-pdf-modul.py --buat-baru
"""
import copy
import html as htmlmod
import importlib.util
import json
import re
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor
from lxml import html as lhtml

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pustaka import nama_berkas_word  # noqa: E402
sys.path.insert(1, str(Path(__file__).resolve().parent.parent))
from svg_word import render_svg  # noqa: E402  (SVG → PNG lewat MuPDF, dipakai bersama TTL)

AKAR = Path(__file__).resolve().parent.parent.parent
KURSUS = AKAR / "Pemodelan-Computer-Aided-Design"
TUJUAN = KURSUS / "Modul-Word"
BASIS = AKAR / "Sistem-Kendali-Cerdas" / "Modul-Word" / "Modul-1-Pengantar-Sistem-Kontrol-Cerdas.docx"
ASESMEN = json.loads((KURSUS / "Attributes" / "Asesmen-Pemodelan-Computer-Aided-Design.json").read_text(encoding="utf-8"))
PAGES = "https://dedik-romahadi.github.io/Mechanical-Engineering-Courses/Pemodelan-Computer-Aided-Design/Modul"
NAMA_MK = ASESMEN["mata_kuliah"]["nama"]
KODE_MK = ASESMEN["mata_kuliah"]["kode_mk"]
SKS = ASESMEN["mata_kuliah"]["sks"]
KELAS = ASESMEN["mata_kuliah"]["kelas_sia"]
# Nama dosen pada sampul, header, dan footer sudah benar di dokumen basis, jadi tidak diganti.
# Dua baris judul sampul (kotak 384,7 pt, Eras Bold ITC) menggantikan
# "SISTEM KENDALI" / "CERDAS" milik dokumen basis.
SAMPUL_BARIS = ("PEMODELAN", "CAD")
# Poin tiap tugas pemodelan T1-T5 mengikuti kartu tugas di halaman modul.
POIN_TUGAS = [6, 6, 6, 11, 11]

# panel "MODUL INTERAKTIF" — pakai pembangun yang sama dengan modul Word course lain
_spec = importlib.util.spec_from_file_location("panel", AKAR / "scripts" / "panel-modul-interaktif-docx.py")
_panel = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_panel)

ARIAL = "Arial"
MONO = "Consolas"
BIRU_H1 = RGBColor(0x1F, 0x38, 0x64)
BIRU_H2 = RGBColor(0x2E, 0x4A, 0x7A)
LEBAR_TEKS_CM = 15.5   # A4 21 cm − margin kiri 3 cm − kanan 2,5 cm

# ─────────────────────────── LaTeX → run teks ───────────────────────────
GREEK = {"alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε", "varepsilon": "ε", "zeta": "ζ", "eta": "η", "theta": "θ",
         "lambda": "λ", "mu": "μ", "nu": "ν", "xi": "ξ", "pi": "π", "rho": "ρ", "sigma": "σ", "tau": "τ", "phi": "φ", "varphi": "φ",
         "chi": "χ", "psi": "ψ", "omega": "ω", "Delta": "Δ", "Gamma": "Γ", "Theta": "Θ", "Lambda": "Λ", "Sigma": "Σ", "Phi": "Φ",
         "Psi": "Ψ", "Omega": "Ω", "ell": "ℓ", "infty": "∞", "partial": "∂", "nabla": "∇"}
SIMBOL = {"times": "×", "cdot": "·", "le": "≤", "leq": "≤", "ge": "≥", "geq": "≥", "ne": "≠", "neq": "≠", "approx": "≈", "propto": "∝",
          "pm": "±", "mp": "∓", "to": "→", "rightarrow": "→", "Rightarrow": "⇒", "leftrightarrow": "↔", "parallel": "∥", "angle": "∠",
          "circ": "°", "ldots": "…", "cdots": "⋯", "dots": "…", "sum": "Σ", "int": "∫", "prod": "Π", "mid": "|", "lvert": "|", "rvert": "|",
          "quad": "  ", "qquad": "    ", "%": "%", ",": " ", ";": " ", ":": " ", "!": "", " ": " ", "{": "{", "}": "}", "_": "_", "&": "&",
          "\\": "; ", "equiv": "≡", "sim": "~", "star": "★", "bullet": "•", "prime": "′", "degree": "°", "langle": "⟨", "rangle": "⟩", "in": "∈"}
FUNGSI = {"ln", "log", "sin", "cos", "tan", "exp", "max", "min", "arcsin", "arccos", "arctan", "sinh", "cosh", "tanh", "lim", "det", "Re", "Im"}


def _baca_grup(s, i):
    """s[i] == '{' → (isi, indeks setelah '}')."""
    kedalaman, j = 0, i
    while j < len(s):
        if s[j] == "{":
            kedalaman += 1
        elif s[j] == "}":
            kedalaman -= 1
            if kedalaman == 0:
                return s[i + 1:j], j + 1
        j += 1
    return s[i + 1:], len(s)


def _argumen(s, i):
    """Argumen setelah perintah: grup {…}, perintah \\x, atau satu karakter."""
    while i < len(s) and s[i] == " ":
        i += 1
    if i >= len(s):
        return "", i
    if s[i] == "{":
        return _baca_grup(s, i)
    if s[i] == "\\":
        m = re.match(r"\\[A-Za-z]+", s[i:])
        if m:
            return m.group(0), i + len(m.group(0))
        return s[i:i + 2], i + 2
    return s[i], i + 1


def latex_runs(s, skrip=None):
    """LaTeX (subset KaTeX yang dipakai modul) → daftar run (teks, skrip) dengan skrip ∈ {None,'sub','sup'}."""
    out = []
    i = 0
    s = s.replace("\\dfrac", "\\frac").replace("\\tfrac", "\\frac").replace("{,}", ",").replace("\\left", "").replace("\\right", "")
    s = s.replace("\\\\", "⁣ROW⁣")   # pemisah baris matriks/cases

    def tambah(teks, sk=skrip):
        if teks:
            if out and out[-1][1] == sk:
                out[-1] = (out[-1][0] + teks, sk)
            else:
                out.append((teks, sk))

    def tambah_runs(rs):
        for teks, sk in rs:
            tambah(teks, sk if sk else skrip)

    while i < len(s):
        c = s[i]
        if c == "\\":
            m = re.match(r"\\([A-Za-z]+)", s[i:])
            if m:
                nama = m.group(1)
                i += len(m.group(0))
                if nama == "frac":
                    a, i = _argumen(s, i)
                    b, i = _argumen(s, i)
                    ra, rb = latex_runs(a), latex_runs(b)
                    ta, tb = "".join(t for t, _ in ra), "".join(t for t, _ in rb)
                    kurung_a = any(ch in ta for ch in "+−- ") and len(ta) > 1
                    kurung_b = any(ch in tb for ch in "+−- ·×/") and len(tb) > 1
                    if kurung_a:
                        tambah("(")
                    tambah_runs(ra)
                    if kurung_a:
                        tambah(")")
                    tambah("/")
                    if kurung_b:
                        tambah("(")
                    tambah_runs(rb)
                    if kurung_b:
                        tambah(")")
                elif nama == "sqrt":
                    if i < len(s) and s[i] == "[":
                        j = s.index("]", i)
                        akar = s[i + 1:j]
                        i = j + 1
                        tambah_runs(latex_runs(akar, "sup"))
                    a, i = _argumen(s, i)
                    ra = latex_runs(a)
                    ta = "".join(t for t, _ in ra)
                    if len(ra) == 1 and re.fullmatch(r"[A-Za-z0-9.,]+", ta):
                        tambah("√" + ta)
                    else:
                        tambah("√(")
                        tambah_runs(ra)
                        tambah(")")
                elif nama in ("text", "mathrm", "mathbf", "mathit", "boldsymbol", "operatorname", "textbf"):
                    a, i = _argumen(s, i)
                    tambah(a.replace("\\ ", " ") if nama == "text" else "".join(t for t, _ in latex_runs(a)))
                elif nama in ("bar", "overline", "hat", "vec", "dot", "ddot", "tilde"):
                    a, i = _argumen(s, i)
                    aksen = {"bar": "̄", "overline": "̅", "hat": "̂", "vec": "⃗", "dot": "̇", "ddot": "̈", "tilde": "̃"}[nama]
                    ra = latex_runs(a)
                    if ra:
                        teks, sk = ra[0]
                        ra[0] = (teks[0] + aksen + teks[1:], sk)
                    tambah_runs(ra)
                elif nama == "begin":
                    a, i = _argumen(s, i)      # bmatrix / pmatrix / cases
                    j = s.find("\\end", i)
                    isi = s[i:j] if j >= 0 else s[i:]
                    i = j if j < 0 else j
                    if j >= 0:
                        _, i = _argumen(s, i + 4)
                    baris = [b.strip() for b in isi.split("⁣ROW⁣") if b.strip()]
                    tambah("[" if "matrix" in a else "{")
                    for k, b in enumerate(baris):
                        if k:
                            tambah("; ")
                        sel = [x.strip() for x in b.split("&")]
                        for m2, x in enumerate(sel):
                            if m2:
                                tambah("  ")
                            tambah_runs(latex_runs(x))
                    tambah("]" if "matrix" in a else "}")
                elif nama in GREEK:
                    tambah(GREEK[nama])
                    if i < len(s) and s[i] == " " and i + 1 < len(s) and (s[i + 1].isalnum() or s[i + 1] in "(|"):
                        i += 1          # \Delta V → ΔV
                elif nama in SIMBOL:
                    tambah(SIMBOL[nama])
                elif nama in FUNGSI:
                    akhir = out[-1][0][-1:] if out else ""
                    tambah((" " if akhir.isalnum() or akhir in ")]" else "") + nama + " ")
                elif nama == "sqrt":
                    pass
                else:
                    tambah(nama)
            else:
                if i + 1 < len(s) and s[i + 1] in SIMBOL:
                    tambah(SIMBOL[s[i + 1]])
                elif i + 1 < len(s):
                    tambah(s[i + 1])
                i += 2
        elif c == "^":
            a, i = _argumen(s, i + 1)
            tambah_runs(latex_runs(a, "sup"))
        elif c == "_":
            a, i = _argumen(s, i + 1)
            tambah_runs(latex_runs(a, "sub"))
        elif c == "{":
            a, i = _baca_grup(s, i)
            tambah_runs(latex_runs(a))
        elif c == "}":
            i += 1
        elif c == "-":
            tambah("−")
            i += 1
        elif c == "~":
            tambah(" ")
            i += 1
        else:
            tambah(c)
            i += 1
    # rapikan spasi ganda; sisa penanda baris (di luar matriks) menjadi ';'
    return [(re.sub(r"[ ]{3,}", "  ", t).replace("⁣ROW⁣", "; "), sk) for t, sk in out if t]


def latex_teks(s):
    return "".join(t for t, _ in latex_runs(s))


# ─────────────────────────── HTML → run teks ───────────────────────────
def _pisah_latex(teks):
    """Teks dengan \\( … \\) → daftar run (teks, bold, italic, skrip)."""
    out = []
    for k, bagian in enumerate(re.split(r"\\\((.*?)\\\)", teks, flags=re.S)):
        if k % 2 == 1:
            out += [(t, False, False, sk) for t, sk in latex_runs(bagian)]
        elif bagian:
            out.append((bagian, False, False, None))
    return out


def runs_dari(el, bold=False, italic=False, skrip=None):
    """Isi campuran elemen HTML → daftar run (teks, bold, italic, skrip)."""
    out = []
    if el.text:
        out += [(t, b or bold, i or italic, sk or skrip) for t, b, i, sk in _pisah_latex(el.text)]
    for anak in el:
        tag = anak.tag if isinstance(anak.tag, str) else ""
        if tag == "br":
            out.append(("\n", bold, italic, skrip))
        elif tag in ("strong", "b"):
            out += runs_dari(anak, True, italic, skrip)
        elif tag in ("em", "i"):
            out += runs_dari(anak, bold, True, skrip)
        elif tag == "sub":
            out += runs_dari(anak, bold, italic, "sub")
        elif tag == "sup":
            out += runs_dari(anak, bold, italic, "sup")
        elif tag == "code":
            out += runs_dari(anak, bold, italic, skrip)
        elif tag in ("script", "style", "svg", "canvas", "button", "input"):
            pass
        else:
            out += runs_dari(anak, bold, italic, skrip)
        if anak.tail:
            out += [(t, b or bold, i or italic, sk or skrip) for t, b, i, sk in _pisah_latex(anak.tail)]
    return out


def rapikan_runs(runs):
    """Gabungkan run bertetangga yang formatnya sama, rapikan spasi."""
    hasil = []
    for teks, b, i, sk in runs:
        teks = re.sub(r"[ \t\r\f\v]+", " ", teks.replace("\xa0", " "))
        if "\n" in teks:
            teks = re.sub(r" *\n *", "\n", teks)
        if not teks:
            continue
        if hasil and hasil[-1][1:] == (b, i, sk):
            hasil[-1] = (hasil[-1][0] + teks, b, i, sk)
        else:
            hasil.append((teks, b, i, sk))
    if hasil:
        hasil[0] = (hasil[0][0].lstrip(), *hasil[0][1:])
        hasil[-1] = (hasil[-1][0].rstrip(), *hasil[-1][1:])
    return [r for r in hasil if r[0]]


def teks_polos(el):
    return "".join(t for t, *_ in rapikan_runs(runs_dari(el))).replace("\n", " ").strip()


# ─────────────────────────── penulisan Word ───────────────────────────
class Penulis:
    def __init__(self, doc):
        self.doc = doc
        self.n_tabel = 0

    def _font(self, run, ukuran=11, bold=False, italic=False, warna=None, nama=ARIAL):
        run.font.name = nama
        run._element.rPr.rFonts.set(qn("w:eastAsia"), nama)
        run._element.rPr.rFonts.set(qn("w:cs"), nama)
        run.font.size = Pt(ukuran)
        run.font.bold = bold
        run.font.italic = italic
        if warna is not None:
            run.font.color.rgb = warna

    def _isi_runs(self, p, runs, ukuran=11, warna=None):
        for teks, b, i, sk in rapikan_runs(runs):
            bagian = teks.split("\n")
            for k, potongan in enumerate(bagian):
                if k:
                    p.add_run().add_break()
                if not potongan:
                    continue
                r = p.add_run(potongan)
                self._font(r, ukuran, b, i, warna)
                if sk == "sub":
                    r.font.subscript = True
                elif sk == "sup":
                    r.font.superscript = True

    def heading1(self, teks):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.15
        self._font(p.add_run(teks.upper()), 14, True, warna=BIRU_H1)
        return p

    def heading2(self, teks):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.line_spacing = 1.15
        self._font(p.add_run(teks), 12, True, warna=BIRU_H2)
        return p

    def para(self, runs, rata=WD_ALIGN_PARAGRAPH.JUSTIFY, ukuran=11, sesudah=6):
        if isinstance(runs, str):
            runs = [(runs, False, False, None)]
        p = self.doc.add_paragraph()
        p.alignment = rata
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(sesudah)
        self._isi_runs(p, runs, ukuran)
        return p

    def bullet(self, runs, ukuran=11):
        if isinstance(runs, str):
            runs = [(runs, False, False, None)]
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(3)
        self._font(p.add_run("•  "), ukuran)
        self._isi_runs(p, runs, ukuran)
        return p

    def persamaan(self, runs, nomor):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_together = True
        p.paragraph_format.tab_stops.add_tab_stop(Cm(LEBAR_TEKS_CM), WD_TAB_ALIGNMENT.RIGHT)
        self._isi_runs(p, [(t, False, False, sk) for t, sk in runs], 11)
        self._font(p.add_run(f"\t({nomor})"), 11)
        return p

    def caption(self, teks):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.15
        if isinstance(teks, str):
            teks = [(teks, True, False, None)]
        self._isi_runs(p, [(t, True, i, sk) for t, b, i, sk in teks], 10)
        return p

    def gambar(self, png, lebar_cm=15.0):
        self.doc.add_picture(str(png), width=Cm(lebar_cm))
        p = self.doc.paragraphs[-1]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        return p

    def tabel(self, header, baris, keterangan):
        self.n_tabel += 1
        self.caption(f"Tabel {self.n_tabel}. {keterangan}")
        self.doc.paragraphs[-1].paragraph_format.keep_with_next = True
        t = self.doc.add_table(rows=1 + len(baris), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(header):
            sel = t.rows[0].cells[j]
            sel._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="DCE6F1"/>'))
            p = sel.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            self._isi_runs(p, [(t_, True, i_, sk) for t_, _, i_, sk in h], 9)
        for i, r in enumerate(baris, 1):
            for j, c in enumerate(r):
                p = t.rows[i].cells[j].paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                self._isi_runs(p, c, 9)
        spasi = self.doc.add_paragraph()
        spasi.paragraph_format.space_after = Pt(4)
        return t

    def kode(self, judul, teks):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        self._font(p.add_run(judul), 10, True, warna=BIRU_H2)
        p = self.doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="6" w:color="CBD5E1"/><w:left w:val="single" w:sz="4" w:space="6" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="4" w:space="6" w:color="CBD5E1"/><w:right w:val="single" w:sz="4" w:space="6" w:color="CBD5E1"/></w:pBdr>'))
        pPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="F1F5F9"/>'))
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.1
        baris = teks.rstrip("\n").split("\n")
        for k, b in enumerate(baris):
            if k:
                p.add_run().add_break()
            if b:
                self._font(p.add_run(b), 8.5, nama=MONO)
        return p

    def kotak(self, runs):
        """Kotak catatan: paragraf beraksen kiri dengan arsiran lembut."""
        p = self.doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        pPr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="2E4A7A"/></w:pBdr>'))
        pPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="F3F6FB"/>'))
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self._isi_runs(p, runs, 10.5)
        return p


# ─────────────────────────── sumber isi ───────────────────────────
IKON_CHIP = re.compile(r"^[\U0001F000-\U0001FAFF←-⯿️\s]+")


def pertemuan(n):
    """Nomor pertemuan modul ke-n. Pertemuan 8 adalah UTS, jadi Modul 8-14 bergeser satu."""
    return n if n <= 7 else n + 1


def rencana_banner(n):
    """Bahan kajian modul ke-n, diambil dari banner LMS pertemuannya.

    Pemodelan CAD tidak punya berkas RPS JavaScript seperti Teknik Tenaga
    Listrik, sehingga "Bahan kajian" diturunkan dari sumber resmi yang ada:
    `Banner/Banner-Pertemuan-P.html` yang dipasang dosen di LMS. Tiap banner
    memuat judul topik (h2), satu paragraf cakupan pertemuan (p sesudahnya),
    dan deretan chip kata kunci. Ketiganya dipakai apa adanya.
    """
    berkas = KURSUS / "Banner" / f"Banner-Pertemuan-{pertemuan(n)}.html"
    root = lhtml.fromstring(berkas.read_text(encoding="utf-8"))
    h2 = root.xpath("//h2")[0]
    cakupan = teks_polos(h2.getparent().xpath("./p")[0])
    # Sebagian chip diawali ikon (🛠, 📐, 💾, ◯); ikonnya hiasan banner, dibuang
    # agar daftar kata kunci seragam di keempat belas dokumen.
    chips = [IKON_CHIP.sub("", teks_polos(c)) for c in root.xpath('//span[contains(@style,"border-radius:7px")]')]
    return teks_polos(h2), cakupan, chips


def indikator(M):
    """Indikator penilaian modul, dirakit dari butir yang benar-benar dinilai.

    Bukan kalimat karangan: jumlah dan bobot soal pilihan ganda, label kelima
    tugas pemodelan beserta poinnya, dan jumlah pertanyaan forum semuanya
    dibaca dari halaman modul yang sedang diproses.
    """
    T = M["tugas"]
    model = "; ".join(f"T{k} {lab} ({POIN_TUGAS[k - 1]} poin)" for k, lab in enumerate(T["model"], 1))
    return (f"menjawab benar {len(T['mc'])} soal pilihan ganda (@1 poin); "
            f"menyerahkan {len(T['model'])} model FreeCAD — {model} — masing-masing berupa berkas .FCStd "
            f"yang diunggah beserta satu angka bacaan geometri yang diperiksa server; "
            f"menulis jawaban {len(M['forum']['tanya'])} pertanyaan Forum Diskusi (minimal 30 kata tiap jawaban).")


def _label_tugas(mentah, nama_array):
    """Label tugas pemodelan dari `const compEzDefs`/`compHardDefs` halaman modul.

    Dibaca per array, bukan sekali jalan atas seluruh berkas, karena penomoran
    CAD (c1-c3 lalu c4-c5) tidak dapat dipisah hanya dari id-nya.
    """
    blok = re.search(rf"const {nama_array} = \[(.*?)\];", mentah, re.S)
    return re.findall(r"\{ id:'c\d+', label:'([^']*)', q:'' \}", blok.group(1)) if blok else []


def muat_modul(n):
    """Baca Modul-N.html → struktur isi."""
    path = KURSUS / "Modul" / f"Modul-{n}.html"
    mentah = path.read_text(encoding="utf-8")
    root = lhtml.fromstring(mentah)
    judul = re.search(r"<title>Modul \d+ — (.*?) \| ", mentah).group(1)
    hero = root.xpath('//div[contains(@class,"hero") and @data-tab="modul"]//p[@class="hero-sub"]')[0]
    svgs = re.findall(r"<figure class=\"ilustrasi[^\"]*\">\s*(<svg.*?</svg>)", mentah, re.S)
    bagian = []
    for sec in root.xpath('//div[@id="page-modul"]//div[@class="section" and @id]') or root.xpath('//div[@class="section" and @id]'):
        label = sec.xpath('.//div[contains(@class,"section-label")]')
        judul_sec = sec.xpath('./h2[contains(@class,"section-title")]')
        if not judul_sec:
            continue
        nama_label = label[0].text_content().strip() if label else ""
        bagian.append({"id": sec.get("id"), "label": nama_label, "judul": teks_polos(judul_sec[0]),
                       "desc": sec.xpath('./p[contains(@class,"section-desc")]'), "el": sec})
    tugas = {"mc": [], "model": _label_tugas(mentah, "compEzDefs") + _label_tugas(mentah, "compHardDefs"),
             "acuan": [], "petunjuk": None}
    for kartu in root.xpath('//div[@id="page-tugas"]//div[@class="mc-card reveal"]'):
        q = kartu.xpath('.//div[@class="mc-q"]')
        opsi = kartu.xpath('.//div[@class="radio-option"]')
        if q:
            tugas["mc"].append((runs_dari(q[0]), [teks_polos(o) for o in opsi]))
    # Kartu tugas: teks soalnya dirakit server per NIM (hanya placeholder di HTML),
    # jadi yang dibawa ke Word adalah gambar acuan beserta keterangannya. SVG diambil
    # dari teks mentah supaya nama atribut kamel (viewBox) tidak dikecilkan parser HTML.
    svg_acuan = dict(re.findall(r'<div class="tugas-gambar" id="gambar-(c\d+)">\s*(<svg.*?</svg>)', mentah, re.S))
    for kartu in root.xpath('//div[@id="page-tugas"]//div[contains(@class,"comp-card")]'):
        cid = (kartu.get("id") or "").replace("card-", "")
        ket = kartu.xpath('.//div[@class="tugas-gambar-ket"]')
        tugas["acuan"].append({"id": cid, "svg": svg_acuan.get(cid),
                               "ket": teks_polos(ket[0]) if ket else ""})
    petunjuk = root.xpath('//div[@id="page-tugas"]//div[contains(@class,"warn-box")]')
    if petunjuk:
        judul_p = petunjuk[0].xpath(".//h4")
        isi_p = petunjuk[0].xpath(".//p")
        tugas["petunjuk"] = (teks_polos(judul_p[0]) if judul_p else "Petunjuk Pengerjaan Tugas",
                             runs_dari(isi_p[0]) if isi_p else [])
    forum = {"skenario": [runs_dari(p) for p in root.xpath('//div[@id="page-forum"]//div[contains(@class,"forum-scenario")]/p')],
             "tanya": [teks_polos(h) for h in root.xpath('//div[@id="page-forum"]//div[@class="fq-head"]/h3')]}
    return {"n": n, "judul": judul, "hero": hero, "svgs": svgs, "bagian": bagian, "tugas": tugas, "forum": forum, "root": root}


# ─────────────────────────── pembangunan dokumen ───────────────────────────
def siapkan_basis(n, judul, abstrak, sub_kode, sub_desk, tmp):
    """Salin basis, ganti identitas sampul/header/footer di XML, hapus badan."""
    z = zipfile.ZipFile(BASIS)
    isi = {nm: z.read(nm) for nm in z.namelist()}
    info = {i.filename: i for i in z.infolist()}
    z.close()
    doc = isi["word/document.xml"].decode("utf8")

    def ganti_t(lama, baru, wajib=True):
        nonlocal doc
        pola = f">{htmlmod.escape(lama, quote=False)}<"
        if wajib and pola not in doc:
            raise SystemExit(f"basis: teks sampul '{lama[:40]}' tidak ditemukan")
        doc = doc.replace(pola, f">{htmlmod.escape(baru, quote=False)}<")

    ganti_t("Modul 1", f"Modul {n}")
    ganti_t("Sub-CPMK 1.1 - Mampu memahami definisi, sejarah dan sistem konfigurasi dari sistem kontrol", f"Sub-CPMK {sub_kode} - {sub_desk}")
    abstrak_lama = re.search(r">(Pertemuan ini membangun kosakata dasar[^<]*)<", doc).group(1)
    doc = doc.replace(f">{abstrak_lama}<", f">{htmlmod.escape(abstrak, quote=False)}<")
    ganti_t("Pengantar Sistem Kontrol Cerdas", judul)
    ganti_t("SISTEM KENDALI", SAMPUL_BARIS[0])
    ganti_t("CERDAS", SAMPUL_BARIS[1])
    ganti_t("Kode Mata kuliah W132500026", f"Kode Mata kuliah {KODE_MK}")
    doc = _atur_sampul(doc, abstrak, judul)
    isi["word/document.xml"] = doc.encode("utf8")
    hdr = isi["word/header1.xml"].decode("utf8").replace("| Sistem Kendali Cerdas", f"| {NAMA_MK}")
    isi["word/header1.xml"] = hdr.encode("utf8")
    ftr = isi["word/footer1.xml"].decode("utf8")
    ftr = re.sub(r"(2025 - 20</w:t>(?:(?!</w:p>).)*?<w:t[^>]*>2</w:t>(?:(?!</w:p>).)*?<w:t[^>]*>)6(</w:t>)", r"\g<1>7\2", ftr, flags=re.S)
    ftr = ftr.replace("2025 - 20<", "2026 - 20<")
    isi["word/footer1.xml"] = ftr.encode("utf8")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as keluar:
        for nm in info:
            keluar.writestr(info[nm], isi[nm])


def _ubah_kotak(doc, kunci, cx=None, cy=None, lebar_pt=None, tinggi_pt=None, sz=None):
    """Ubah geometri kotak teks sampul yang memuat `kunci` (wps + VML fallback)."""
    pos = [m.start() for m in re.finditer(re.escape(kunci), doc)]
    if len(pos) < 2:
        raise SystemExit(f"sampul: kotak '{kunci[:30]}' tidak ditemukan dua kali (wps + VML)")
    # wps (kemunculan pertama): wp:extent + a:ext di dalam anchor
    a = doc.rfind("<wp:anchor", 0, pos[0]); b = doc.find("</wp:anchor>", pos[0])
    blok = doc[a:b]
    if cx:
        blok = re.sub(r'(<wp:extent cx=")\d+', lambda m: m.group(1) + str(cx), blok, count=1)
        blok = re.sub(r'(<a:ext cx=")\d+', lambda m: m.group(1) + str(cx), blok, count=1)
    if cy:
        blok = re.sub(r'(<wp:extent cx="\d+" cy=")\d+', lambda m: m.group(1) + str(cy), blok, count=1)
        blok = re.sub(r'(<a:ext cx="\d+" cy=")\d+', lambda m: m.group(1) + str(cy), blok, count=1)
    if sz:
        blok = re.sub(r'(<w:sz w:val=")\d+("/>\s*<w:szCs w:val=")\d+', lambda m: f"{m.group(1)}{sz}{m.group(2)}{sz}", blok)
    doc = doc[:a] + blok + doc[b:]
    # VML fallback (kemunculan kedua)
    pos = [m.start() for m in re.finditer(re.escape(kunci), doc)]
    v = doc.rfind("<v:shape", 0, pos[1]); w = doc.find("</v:shape>", pos[1])
    blok = doc[v:w]
    if lebar_pt:
        blok = re.sub(r"width:[\d.]+pt", f"width:{lebar_pt}pt", blok, count=1)
    if tinggi_pt:
        blok = re.sub(r"height:[\d.]+pt", f"height:{tinggi_pt}pt", blok, count=1)
    if sz:
        blok = re.sub(r'(<w:sz w:val=")\d+("/>\s*<w:szCs w:val=")\d+', lambda m: f"{m.group(1)}{sz}{m.group(2)}{sz}", blok)
    return doc[:v] + blok + doc[w:]


def _atur_sampul(doc, abstrak, judul):
    doc = _ubah_kotak(doc, ">Modul ", cx=1750000, lebar_pt=137.8)
    # kotak judul: lebar penuh dan dua baris agar judul panjang (Modul 4, 6, 12) tidak terpotong
    doc = _ubah_kotak(doc, f">{htmlmod.escape(judul, quote=False)}<", cx=6400000, cy=800000, lebar_pt=504, tinggi_pt=63)
    doc = _ubah_kotak(doc, htmlmod.escape(abstrak, quote=False)[:40], cy=730000, tinggi_pt=57.5, sz=18)
    doc = _ubah_kotak(doc, ">Sub-CPMK ", cy=730000, tinggi_pt=57.5, sz=18)
    return doc


def bersihkan_media(path):
    """Buang relasi dan media dokumen yang tidak lagi dirujuk (gambar Sisken lama)."""
    z = zipfile.ZipFile(path)
    isi = {nm: z.read(nm) for nm in z.namelist()}
    info = {i.filename: i for i in z.infolist()}
    z.close()
    doc = isi["word/document.xml"].decode("utf8")
    rels = isi["word/_rels/document.xml.rels"].decode("utf8")
    dipakai = set(re.findall(r'r:(?:embed|id|link)="(rId\d+)"', doc))
    buang_media = []

    def saring(m):
        rid, target = m.group(1), m.group(2)
        if "media/" in target and rid not in dipakai:
            buang_media.append("word/" + target)
            return ""
        return m.group(0)
    rels = re.sub(r'<Relationship [^>]*Id="(rId\d+)"[^>]*Target="([^"]+)"[^>]*/>', saring, rels)
    lain = set()
    for nm, b in isi.items():
        if nm.endswith(".rels") and nm != "word/_rels/document.xml.rels":
            lain |= {"word/" + t for t in re.findall(r'Target="(media/[^"]+)"', b.decode("utf8"))}
    isi["word/_rels/document.xml.rels"] = rels.encode("utf8")
    for m in buang_media:
        if m in isi and m not in lain:
            del isi[m]
            del info[m]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as keluar:
        for nm in info:
            keluar.writestr(info[nm], isi[nm])


def bangun(n, tmpdir):
    M = muat_modul(n)
    sub = ASESMEN["sub_cpmk"][n - 1]
    sub_kode = sub["kode"].replace("Sub-CPMK ", "")
    P = pertemuan(n)
    hero_teks = teks_polos(M["hero"])
    # abstrak sampul: ≤ ~165 karakter (4 baris kotak sampul) pada batas kalimat
    abstrak = ""
    for kal in re.split(r"(?<=[.!?])\s+", hero_teks):
        if len(abstrak) + len(kal) > 165 and abstrak:
            break
        abstrak += (" " if abstrak else "") + kal
    if len(abstrak) > 170:      # kalimat pertama terlalu panjang: potong pada batas klausa
        potong = abstrak[:160]
        k = max(potong.rfind(", "), potong.rfind("; "), potong.rfind(": "))
        abstrak = (potong[:k] if k > 90 else potong.rsplit(" ", 1)[0]).rstrip(",;: ") + "."
    nama = nama_berkas_word(n, M["judul"])
    url = f"{PAGES}/Modul-{n}.html"
    tmp = Path(tmpdir) / f"basis-{n}.docx"
    siapkan_basis(n, M["judul"], abstrak, sub_kode, sub["deskripsi"], tmp)

    doc = Document(str(tmp))
    body = doc.element.body
    anak = list(body)
    for el in anak[3:]:
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    # panel MODUL INTERAKTIF (bentuk sama dengan course lain)
    ekor = ("Pada layar masuk, pilih tombol “👁️ Mode Preview (Tanpa Login)” untuk membaca seluruh materi tanpa perlu NIM dan PIN. "
            "Untuk mengerjakan Tugas dan Forum, masuk sebagai Mahasiswa memakai NIM dan PIN.")
    rid = doc.part.relate_to(url, RT.HYPERLINK, is_external=True)
    panel = parse_xml(f'<w:body {nsdecls("w", "r")}>{_panel.bangun(url, ekor, rid)}</w:body>')
    p1, p2 = list(panel)
    anak[2].addnext(p2)
    anak[2].addnext(p1)

    W = Penulis(doc)
    # ── PENDAHULUAN ──
    W.heading1("Pendahuluan")
    W.para(runs_dari(M["hero"]))
    _, cakupan, chips = rencana_banner(n)
    W.heading2("Capaian Pembelajaran (Sub-CPMK)")
    W.bullet([(f"Sub-CPMK {sub_kode}: ", True, False, None), (sub["deskripsi"], False, False, None)])
    W.bullet([("Bahan kajian: ", True, False, None), (cakupan, False, False, None)])
    W.bullet([("Kata kunci: ", True, False, None), ("; ".join(chips) + ".", False, False, None)])
    W.bullet([("Indikator: ", True, False, None), (indikator(M), False, False, None)])
    W.bullet([("Kedudukan: ", True, False, None), (f"Pertemuan {P} dari 16, mata kuliah {SKS} SKS kelas {KELAS}; bobot penilaian Tugas {sub['bobot']['tugas']} %, UTS {sub['bobot']['uts']} %, UAS {sub['bobot']['uas']} % (SIA).", False, False, None)])

    # ── BAGIAN MATERI ──
    n_gambar = 0
    pustaka = []
    for B in M["bagian"]:
        if B["id"] == "m-pustaka":
            for kartu in B["el"].xpath('.//div[contains(@class,"reference-card")]'):
                pustaka.append(rapikan_runs(runs_dari(kartu.xpath("./div")[0])))
            continue
        W.heading1(B["judul"])
        for d in B["desc"]:
            W.para(runs_dari(d))
        for el in B["el"]:
            tag = el.tag if isinstance(el.tag, str) else ""
            kelas = el.get("class", "") or ""
            if tag == "figure":
                svg = M["svgs"][n_gambar] if n_gambar < len(M["svgs"]) else None
                n_gambar += 1
                cap = el.xpath("./figcaption")
                if svg:
                    png = Path(tmpdir) / f"m{n}-g{n_gambar}.png"
                    render_svg(svg, png)
                    W.gambar(png)
                if cap:
                    W.caption([(t, b, i, sk) for t, b, i, sk in rapikan_runs(runs_dari(cap[0]))])
            elif "formula-block" in kelas:
                label = el.xpath('./div[@class="formula-label"]')
                utama = el.xpath('./div[@class="formula-main"]')
                nomor = el.xpath('.//span[@class="formula-number"]')
                desc = el.xpath('./div[@class="formula-desc"]')
                if label:
                    W.heading2(teks_polos(label[0]))
                if utama:
                    m = re.search(r"\\\((.*)\\\)", (utama[0].text or "") + "".join(lhtml.tostring(c, encoding="unicode") for c in utama[0]), re.S)
                    latex = m.group(1) if m else teks_polos(utama[0])
                    no = teks_polos(nomor[0]).strip("()") if nomor else ""
                    W.persamaan(latex_runs(latex), no)
                if desc:
                    W.para(runs_dari(desc[0]))
            elif "rumus-jelas" in kelas:
                daftar = el.xpath('.//div[contains(@class,"anim-var-list")]')
                salinan = copy.deepcopy(el)
                for d in salinan.xpath('.//div[contains(@class,"anim-var-list")]'):
                    d.getparent().remove(d)
                W.para(runs_dari(salinan))
                for d in daftar:
                    for span in d.xpath('./span[contains(@class,"anim-var")]'):
                        notasi = span.xpath('./span[@class="rumus-notasi"]')
                        arti = span.xpath("./span[not(@class)]")
                        r = []
                        if notasi:
                            r += [(t, True, False, sk) for t, sk in latex_runs(re.sub(r"^\\\(|\\\)$", "", (notasi[0].text or "").strip()))]
                        if arti:
                            r += [(": " + teks_polos(arti[0]), False, False, None)]
                        if r:
                            W.bullet(r, 10.5)
            elif "cards" in kelas.split():
                for kartu in el.xpath('./div[@class="card"]'):
                    h3 = kartu.xpath("./h3")
                    p = kartu.xpath("./p")
                    rumus = kartu.xpath('./div[@class="formula"]')
                    r = [(teks_polos(h3[0]) + ": ", True, False, None)] if h3 else []
                    if p:
                        r += runs_dari(p[0])
                    if rumus:
                        r += [(" ", False, False, None)] + [(t, False, True, sk) for t, sk in latex_runs(re.sub(r"\\\(|\\\)", "", teks_polos(rumus[0])))]
                    W.bullet(r)
            elif "tbl-wrap" in kelas:
                header = [rapikan_runs(runs_dari(th)) for th in el.xpath(".//thead//th")]
                baris = [[rapikan_runs(runs_dari(td)) for td in tr.xpath("./td")] for tr in el.xpath(".//tbody/tr")]
                if header:
                    judul_tabel = "".join(t for t, *_ in header[0])
                    W.tabel(header, baris, f"{judul_tabel}: {', '.join(''.join(t for t, *_ in h) for h in header[1:])}.")
            elif "info-box" in kelas or "tip-box" in kelas or "warning-box" in kelas:
                W.kotak(runs_dari(el))
            elif "anim-panel" in kelas:
                judul_anim = el.xpath('.//span[@class="anim-title"]')
                cara = el.xpath('.//div[contains(@class,"tip-box")]')
                if judul_anim:
                    W.heading2(teks_polos(judul_anim[0]))
                W.para("Animasi ini dapat dijalankan pada halaman modul interaktif (tautan di awal dokumen); panel di bawah merangkum cara membacanya.")
                if cara:
                    W.kotak(runs_dari(cara[0]))
            elif "code-wrap" in kelas:
                label = el.xpath('.//span[@class="code-label"]')
                pre = el.xpath(".//pre")
                if pre:
                    W.kode(teks_polos(label[0]) if label else "Kode", htmlmod.unescape(pre[0].text_content()))
            # elemen lain (divider, label, judul, desc) sudah ditangani

    # ── TUGAS & FORUM ──
    T = M["tugas"]
    if T["mc"]:
        total_poin = len(T["mc"]) + sum(POIN_TUGAS[:len(T["model"])])
        W.heading1("Tugas Modul")
        W.para(f"Tugas Modul {n} dikerjakan pada halaman modul interaktif dan dinilai otomatis oleh server: "
               f"{len(T['mc'])} soal pilihan ganda (1 poin) dan {len(T['model'])} tugas pemodelan FreeCAD "
               f"berbobot {'/'.join(str(p) for p in POIN_TUGAS[:len(T['model'])])} poin, total {total_poin} poin. "
               "Tiap tugas pemodelan menuntut dua hal berurutan: berkas .FCStd hasil pemodelan diunggah lebih dulu, "
               "lalu satu angka bacaan geometri (misalnya Area, Shape.Length, atau Shape.Volume) disalin dari FreeCAD "
               "sebagai jawaban. Dimensi tiap tugas diturunkan dari NIM sehingga angka tiap mahasiswa berbeda; "
               "teks tugasnya dirakit server saat mahasiswa masuk dan karena itu tidak dicetak pada modul ini. "
               "Yang dapat disiapkan lebih awal adalah butir pilihan ganda dan gambar acuan tiap tugas berikut.")
        if T["petunjuk"]:
            W.heading2(T["petunjuk"][0])
            W.kotak(T["petunjuk"][1])
        W.heading2("Bagian A — Pilihan Ganda")
        for k, (q, opsi) in enumerate(T["mc"], 1):
            W.para([(f"{k}. ", True, False, None)] + q, sesudah=2)
            for j, o in enumerate(opsi):
                W.bullet([(o, False, False, None)], 10.5)
        W.heading2("Bagian B — Tugas Pemodelan FreeCAD (unggah .FCStd + angka bacaan)")
        for k, lab in enumerate(T["model"], 1):
            W.para([(f"T{k}. ", True, False, None), (lab, False, False, None),
                    (f"  ({POIN_TUGAS[k - 1]} poin)", True, False, None)], sesudah=2)
            acuan = T["acuan"][k - 1] if k - 1 < len(T["acuan"]) else None
            if acuan and acuan["svg"]:
                png = Path(tmpdir) / f"m{n}-t{k}.png"
                render_svg(acuan["svg"], png)
                W.gambar(png, 13.0)
            if acuan and acuan["ket"]:
                W.caption([(acuan["ket"], True, False, None)])
    F = M["forum"]
    if F["tanya"]:
        W.heading1("Forum Diskusi")
        for p in F["skenario"]:
            W.para(p)
        for k, t in enumerate(F["tanya"], 1):
            W.bullet([(f"Pertanyaan {k}: ", True, False, None), (t, False, False, None)])
        W.para("Jawaban ditulis pada halaman modul interaktif (minimal 30 kata per pertanyaan) dan menjadi syarat kelengkapan modul bersama tugas.")

    # ── DAFTAR PUSTAKA ──
    W.heading1("Daftar Pustaka")
    entri = []
    for r in pustaka:
        teks = "".join(t for t, *_ in r).split("\n")[0]
        teks = re.sub(r"^\[\d+\]\s*", "", teks).strip()
        # Nama penulis yang sudah berakhir titik ("… dkk.") bertemu titik pemisah
        # dari templat kartu; di dokumen cetak titik gandanya dirapikan.
        entri.append(re.sub(r"(?<!\.)\.\.(?!\.)", ".", teks))
    for e in sorted(set(entri), key=lambda s: s.lower()):
        p = W.para([(e, False, False, None)], rata=WD_ALIGN_PARAGRAPH.LEFT, ukuran=10.5, sesudah=4)
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        p.paragraph_format.line_spacing = 1.2

    TUJUAN.mkdir(exist_ok=True)
    keluar = TUJUAN / f"{nama}.docx"
    doc.save(str(keluar))
    bersihkan_media(keluar)
    return keluar


def main():
    arg = [a for a in sys.argv[1:] if not a.startswith("--")]
    nomor = list(range(1, 15)) if "--semua" in sys.argv else [int(a) for a in arg]
    if not nomor:
        raise SystemExit(__doc__)
    with tempfile.TemporaryDirectory() as tmpdir:
        for n in nomor:
            keluar = bangun(n, tmpdir)
            print(f"  {keluar.relative_to(AKAR)} ({keluar.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
